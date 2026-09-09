"""Firewall assessment service.

Assessment results are produced either by the live Azure Functions (netsec)
or by the local compliance engine (sample fallback) and cached in memory for
``CACHE_TTL``. Everything that used to be persisted in JSON documents is now
persisted through repositories:

* ``assessment_stats``       - run counter + last run timestamp
* ``assessment_history``     - compliance trend snapshots
* ``findings``               - per-run finding rows (audit / netsec writes)
"""

import json
import logging
import os
import random
import sys
import time

import requests

from config import settings
from database.db import get_session
from database.repositories import AssessmentsRepository
from database.repositories import FindingsRepository
from services import timeutil

log = logging.getLogger("assessment")


class AssessmentUnavailableError(Exception):
    """Raised when a firewall must be reached for a live assessment but is not.

    Carries a human-readable message the UI surfaces to the user so a stopped
    or offline firewall VM produces a clear message instead of a silent hang.
    """


def _would_run_live():
    """True when ``_live_assessment`` would really call the Azure function."""
    if not settings.LIVE_ENABLED:
        return False
    key = (settings.FUNCTION_KEY or "").strip()
    return bool(key and not key.startswith("PLACEHOLDER"))


def _ensure_firewall_reachable(firewall_id):
    """Fast pre-flight before a live assessment.

    Only acts when a live function call would actually be attempted. When the
    target firewall's management host is known and does not answer a short TCP
    probe, raise ``AssessmentUnavailableError`` immediately (instead of waiting
    on the function timeout) with a clear, user-facing message.
    """
    if not _would_run_live():
        return
    from services import managed_firewalls_service

    host_ip = managed_firewalls_service.resolve_device_host(firewall_id)
    if not host_ip:
        return
    if not managed_firewalls_service.probe_host(host_ip, force=True):
        raise AssessmentUnavailableError(
            "{0} is not reachable at {1}. The firewall VM appears to be "
            "stopped or offline. Start the VM and once it is back online, "
            "run the assessment again.".format(firewall_id or "vmpafw01", host_ip)
        )

UI_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))

FUNCTIONS_ROOT = os.path.abspath(
    os.path.join(UI_ROOT, "..", "netsec-agent", "functions")
)

# On the Azure Web App only `ui/` is shipped, so the `netsec-agent` source tree
# is not present. Fall back to a bundled copy of the report/compliance modules
# that the CI workflow copies into the UI package before deployment.
if not os.path.isdir(FUNCTIONS_ROOT):
    FUNCTIONS_ROOT = os.path.join(UI_ROOT, "netsec_functions")

BASELINE_PATH = os.path.join(
    FUNCTIONS_ROOT,
    "baseline",
    "baseline_rules.json"
)

EXCEL_DIR = os.path.join(
    UI_ROOT,
    "static",
    "reports"
)

EXCEL_FILE = os.path.join(
    EXCEL_DIR,
    "PaloAlto_Assessment.xlsx"
)

SEVERITY_KEYS = ("critical", "high", "medium", "low")

if FUNCTIONS_ROOT not in sys.path:
    sys.path.insert(0, FUNCTIONS_ROOT)

SERVICE_DIR = os.path.dirname(os.path.abspath(__file__))
if SERVICE_DIR not in sys.path:
    sys.path.insert(0, SERVICE_DIR)

FIREWALLS = ["vmpafw01", "vmpafw02"]

_cache = {
    "assessment": {},
    "ts": {}
}


def _assessments_repo():
    return AssessmentsRepository(get_session())


# ------------------------------------------------------------
# PERSISTENCE (PostgreSQL via repositories)
# ------------------------------------------------------------

def get_assessment_stats():
    """Return ``{assessments_run, last_assessment_ts}`` from assessment_stats."""
    return _assessments_repo().stats_as_dict()


def get_history():
    """Return stored assessment history snapshots (ascending by time)."""
    return _load_history()


def _snapshot_from_row(row):
    return {
        "run_id": row.assessment_id,
        "ts": row.executed_at,
        "firewall_name": row.firewall_name,
        "compliance_pct": row.compliance_score,
        "security_score": row.security_score,
        "severity": {
            "critical": int(row.critical_findings or 0),
            "high": int(row.high_findings or 0),
            "medium": int(row.medium_findings or 0),
            "low": int(row.low_findings or 0),
        },
        "finding_count": int(row.total_findings or 0),
    }


def _load_history():
    rows = _assessments_repo().list_history()
    snapshots = [_snapshot_from_row(row) for row in rows]
    if not snapshots:
        _seed_history()
        rows = _assessments_repo().list_history()
        snapshots = [_snapshot_from_row(row) for row in rows]
    return snapshots


def _seed_history():
    rng = random.Random(20260814)
    now = time.time()
    rows = []
    base = 36.0
    for i in range(12):
        pct = round(base + rng.uniform(-4.5, 3.0), 1)
        severity = {
            "critical": rng.randint(9, 14),
            "high": rng.randint(9, 14),
            "medium": rng.randint(2, 6),
            "low": rng.randint(0, 3),
        }
        points = (
            severity["critical"] * 4.0
            + severity["high"] * 2.0
            + severity["medium"] * 1.0
            + severity["low"] * 0.5
        )
        rows.append({
            "run_id": "ASM-{0:06d}".format(i + 1),
            "ts": now - (12 - i) * 86400,
            "compliance_pct": pct,
            "security_score": round(
                100.0 * (1.0 - points / (44 * 4.0)),
                1,
            ),
            "severity": severity,
            "finding_count": sum(severity.values()),
        })
    _assessments_repo().seed_history(rows)


def _record_history(snapshot):
    """Persist one trend snapshot (dedupes snapshots < 300s apart)."""
    try:
        _assessments_repo().record(snapshot)
    except Exception as exc:
        log.warning("Assessment history write failed: %s", exc)


def _record_assessment(data=None):
    """Increment the assessment run counter for a fresh run."""
    try:
        stats = _assessments_repo().increment_run()
        return "ASM-{0:06d}".format(int(stats.assessments_run or 0))
    except Exception as exc:
        log.warning("Assessment stats write failed: %s", exc)
        return None


def _persist_findings(run_id, findings, firewall_id):
    """Store the per-run findings rows (audit + netsec-style persistence)."""
    try:
        FindingsRepository(get_session()).replace_for_assessment(
            run_id,
            findings,
            firewall_name=firewall_id or "vmpafw01",
        )
    except Exception as exc:
        log.warning("Findings write failed: %s", exc)


# ------------------------------------------------------------
# DATA PRODUCERS
# ------------------------------------------------------------

def _apply_firewall(data, firewall_id):
    """Label an assessment snapshot with the selected firewall name.

    Both firewalls point at the same underlying firewall system; only the
    logical device name differs (vmpafw01 / vmpafw02).
    """
    inventory = dict(data.get("inventory") or {})
    inventory["hostname"] = firewall_id
    data["inventory"] = inventory
    data["_firewall_id"] = firewall_id
    return data


def _severity_breakdown(findings):
    severity = {"critical": 0, "high": 0, "medium": 0, "low": 0}
    for f in findings:
        r = (f.get("risk") or "").lower()
        if r in severity:
            severity[r] += 1
    return severity


def _live_assessment():

    if not settings.LIVE_ENABLED:
        return None

    key = (settings.FUNCTION_KEY or "").strip()
    if not key or key.startswith("PLACEHOLDER"):
        return None

    try:

        response = requests.get(
            f"{settings.BASE_URL}/run_compliance_assessment",
            params={
                "code": settings.FUNCTION_KEY
            },
            timeout=settings.LIVE_TIMEOUT
        )

        if response.status_code == 200:

            payload = response.json()

            if (
                isinstance(payload, dict)
                and "summary" in payload
            ):

                return payload

        return None

    except Exception:

        return None


def _local_assessment():

    from compliance.compliance_engine import (
        ComplianceEngine
    )

    from compliance.findings_generator import (
        FindingsGenerator
    )

    from sample_assessment import (
        SAMPLE_ASSESSMENT
    )

    with open(
        BASELINE_PATH,
        "r"
    ) as handle:

        baseline = json.load(handle)

    results = (
        ComplianceEngine(baseline)
        .evaluate(SAMPLE_ASSESSMENT)
    )

    findings = (
        FindingsGenerator()
        .generate(results)
    )

    summary = {
        "total_controls":
            len(results),
        "compliant":
            len([
                r
                for r in results
                if r["status"] == "COMPLIANT"
            ]),
        "non_compliant":
            len([
                r
                for r in results
                if r["status"] == "NON_COMPLIANT"
            ]),
        "not_assessed":
            len([
                r
                for r in results
                if r["status"] == "NOT_ASSESSED"
            ])
    }

    data = dict(SAMPLE_ASSESSMENT)

    data["summary"] = summary
    data["findings"] = findings
    data["assessment"] = results

    return data


def _stamp(data, source):

    data["_source"] = source
    data["_collected_at"] = (
        timeutil.ist_now().isoformat()
    )

    return data


def get_full_assessment(firewall_id="vmpafw01", force=False, require_reachable=False):

    firewall_id = firewall_id or "vmpafw01"
    now = time.time()

    if require_reachable:
        _ensure_firewall_reachable(firewall_id)

    cached = _cache["assessment"].get(firewall_id)

    if (
        not force
        and cached
        and now - _cache["ts"].get(firewall_id, 0)
        < settings.CACHE_TTL
    ):

        return cached

    # vmpafw02 mirrors vmpafw01 (same system, different logical name).
    if firewall_id != "vmpafw01" and "vmpafw01" in _cache["assessment"]:
        import copy

        data = copy.deepcopy(_cache["assessment"]["vmpafw01"])
        data = _apply_firewall(data, firewall_id)
        _cache["assessment"][firewall_id] = data
        _cache["ts"][firewall_id] = now
        return data

    data = _live_assessment()

    if data is None:

        data = _local_assessment()

        data = _stamp(
            data,
            "sample"
        )

    else:

        data = _stamp(
            data,
            "live"
        )

    data = _apply_firewall(data, firewall_id)

    run_id = _record_assessment()
    data["_run_id"] = run_id

    _persist_findings(run_id, data.get("findings") or [], firewall_id)

    _cache["assessment"][firewall_id] = data
    _cache["ts"][firewall_id] = now

    return data


def ingest_live_payload(data):
    """Persist a full assessment payload produced outside this process
    (live function call). Refreshes the memory cache and records the run."""
    firewall_id = data.get("_firewall_id") or "vmpafw01"
    data = _apply_firewall(data, firewall_id)

    run_id = _record_assessment()
    data["_run_id"] = run_id

    _persist_findings(run_id, data.get("findings") or [], firewall_id)

    _cache["assessment"][firewall_id] = data
    _cache["ts"][firewall_id] = time.time()
    return data


def get_summary(firewall_id="vmpafw01", force=False):

    data = get_full_assessment(firewall_id, force)

    summary = dict(
        data.get(
            "summary",
            {}
        )
    )

    summary["source"] = (
        data.get("_source", "live")
    )

    summary["collected_at"] = (
        data.get("_collected_at")
    )

    return summary


def get_findings(firewall_id="vmpafw01", force=False):

    data = get_full_assessment(firewall_id, force)

    return {
        "findings":
            data.get(
                "findings",
                []
            ),
        "source":
            data.get(
                "_source",
                "live"
            ),
        "collected_at":
            data.get(
                "_collected_at"
            )
    }


def get_posture(firewall_id="vmpafw01", force=False):
    data = get_full_assessment(firewall_id, force)

    summary = data.get("summary", {})
    findings = data.get("findings", [])
    results = data.get("assessment", [])
    inventory = data.get("inventory", {})

    total = summary.get("total_controls", 0) or len(results)
    compliant = summary.get("compliant", 0)
    non_compliant = summary.get("non_compliant", 0)
    not_assessed = summary.get("not_assessed", 0)
    compliance_pct = round(compliant / total * 100, 1) if total else 0.0

    severity = _severity_breakdown(findings)

    risk_points = (
        severity.get("critical", 0) * 4.0
        + severity.get("high", 0) * 2.0
        + severity.get("medium", 0) * 1.0
        + severity.get("low", 0) * 0.5
    )
    max_points = total * 4.0
    security_score = (
        round(100.0 * (1.0 - risk_points / max_points), 1)
        if max_points
        else 100.0
    )

    collected_at = data.get("_collected_at")
    stats = get_assessment_stats()
    run_id = data.get("_run_id") or "ASM-{0:06d}".format(int(stats.get("assessments_run", 0)))

    prev_snapshots = _load_history()
    prev = None
    for s in reversed(prev_snapshots):
        if s.get("firewall_name", "vmpafw01") == firewall_id:
            prev = s
            break

    trend_pct = 0.0
    severity_change = {k: 0 for k in SEVERITY_KEYS}
    if prev and isinstance(prev.get("compliance_pct"), (int, float)):
        trend_pct = round(compliance_pct - prev["compliance_pct"], 1)
        prev_sev = prev.get("severity") or {}
        for k in SEVERITY_KEYS:
            severity_change[k] = int(severity.get(k, 0)) - int(prev_sev.get(k, 0))

    snapshot = {
        "run_id": run_id,
        "ts": time.time(),
        "firewall_name": firewall_id,
        "compliance_pct": compliance_pct,
        "security_score": security_score,
        "severity": severity,
        "finding_count": len(findings),
    }
    _record_history(snapshot)

    history = _load_history()

    posture = {
        "security_score": security_score,
        "compliance_pct": compliance_pct,
        "trend_pct": trend_pct,
        "compliant": compliant,
        "non_compliant": non_compliant,
        "not_assessed": not_assessed,
        "total_controls": total,
        "severity": severity,
        "severity_change": severity_change,
        "assessments_run": int(stats.get("assessments_run", 0)),
        "run_id": run_id,
        "last_assessment_ts": stats.get("last_assessment_ts"),
        "collected_at": collected_at,
        "_source": data.get("_source", "sample"),
        "firewall_id": firewall_id,
    }

    firewall = {
        "hostname": inventory.get("hostname", firewall_id),
        "model": inventory.get("model", ""),
        "version": inventory.get("version", ""),
        "serial": inventory.get("serial", ""),
    }

    return {
        "posture": posture,
        "firewall": firewall,
        "firewalls": [firewall],
        "findings": findings,
        "assessment": results,
        "history": history,
        "_source": data.get("_source", "sample"),
        "_collected_at": collected_at,
        "_firewall_id": firewall_id,
    }


def get_executive_summary(firewall_id="vmpafw01", force=False):

    from reports.executive_summary import (
        ExecutiveSummary
    )

    data = get_full_assessment(firewall_id, force)

    summary = (
        ExecutiveSummary()
        .generate(data)
    )

    summary["_source"] = (
        data.get("_source", "live")
    )

    summary["_collected_at"] = (
        data.get("_collected_at")
    )

    summary["_firewall_id"] = firewall_id

    return summary


def get_executive_summary_pdf(firewall_id="vmpafw01", force=False):

    from reports.executive_summary_pdf import (
        generate as generate_pdf
    )

    firewall_id = firewall_id or "vmpafw01"
    summary = get_executive_summary(firewall_id, force)

    month = timeutil.ist_now().strftime("%b_%Y")
    filename = "Executive_Summary_{0}_{1}.pdf".format(firewall_id, month)

    os.makedirs(
        EXCEL_DIR,
        exist_ok=True
    )

    output_file = os.path.join(
        EXCEL_DIR,
        filename
    )

    generate_pdf(
        summary,
        output_file
    )

    return {
        "summary": summary,
        "filename": filename,
        "local_file": output_file,
        "download_url": "reports/{0}".format(filename),
    }


def get_excel_report(firewall_id="vmpafw01", force=False):

    from reports.excel_report import (
        ExcelReport
    )

    firewall_id = firewall_id or "vmpafw01"
    data = get_full_assessment(firewall_id, force)

    os.makedirs(
        EXCEL_DIR,
        exist_ok=True
    )

    filename = "{0}_Assessment_Workbook.xlsx".format(firewall_id)
    output_file = os.path.join(
        EXCEL_DIR,
        filename
    )

    (
        ExcelReport()
        .generate(
            data,
            output_file=output_file
        )
    )

    return {
        "status":
            "SUCCESS",
        "message":
            "Excel report generated successfully.",
        "summary":
            data.get(
                "summary",
                {}
            ),
        "local_file":
            output_file,
        "download_url":
            "reports/{0}".format(filename),
        "_source":
            data.get(
                "_source",
                "live"
            ),
        "firewall_id": firewall_id,
    }


# ------------------------------------------------------------
# ESTATE (FULL INVENTORY) ASSESSMENT
# ------------------------------------------------------------

def _pct(n, total):
    return round(n / total * 100.0, 1) if total else 0.0


def _device_payload(entry, data):
    """Summarise one managed device (registry row + assessment snapshot).

    Clones share the underlying system with their parent (``clone_of``), so a
    clone's snapshot is the parent's data relabelled with the clone's name.
    """
    summary = data.get("summary") or {}
    results = data.get("assessment") or []
    findings = data.get("findings") or []
    total = summary.get("total_controls", 0) or len(results)
    compliant = int(summary.get("compliant", 0))
    non_compliant = int(summary.get("non_compliant", 0))
    not_assessed = int(summary.get("not_assessed", 0))
    return {
        "id": entry.get("id"),
        "device": entry.get("device_name"),
        "host_ip": entry.get("host_ip"),
        "status": entry.get("status") or "down",
        "clone_of": entry.get("clone_of"),
        "is_clone": bool(entry.get("clone_of")),
        "compliance_pct": _pct(compliant, total),
        "compliant": compliant,
        "non_compliant": non_compliant,
        "not_assessed": not_assessed,
        "total_controls": total,
        "findings_count": len(findings),
        "severity": _severity_breakdown(findings),
        "source": data.get("_source", "sample"),
        "collected_at": data.get("_collected_at"),
        "summary": summary,
        "assessment": results,
        "findings": findings,
    }


def get_estate_assessment(force=False):
    """Assess every managed firewall in the inventory together.

    Returns cumulative compliance statistics (average compliant / non
    compliant / not assessed percentages across all registered devices)
    plus a per-device payload that includes each device's control rows and
    findings, so reports and the Security Operations Centre can render the
    full estate.

    The estate is computed from a single underlying assessment snapshot that
    is then relabelled per registered device (originals and clones).
    """
    from services import managed_firewalls_service

    entries = managed_firewalls_service.list_firewalls()
    devices = []
    if entries:
        names = [entry["device_name"] for entry in entries]
        prime = "vmpafw01" if "vmpafw01" in names else names[0]
        base = get_full_assessment(prime, force=force)
        import copy

        for entry in entries:
            name = entry.get("device_name")
            if name == prime:
                data = base
            else:
                data = copy.deepcopy(base)
                data = _apply_firewall(data, name)
                _cache["assessment"][name] = data
                _cache["ts"][name] = time.time()
            devices.append(_device_payload(entry, data))

    device_count = len(devices)
    cumulative = {
        "device_count": device_count,
        "total_controls": sum(d["total_controls"] for d in devices),
        "total_compliant": sum(d["compliant"] for d in devices),
        "total_non_compliant": sum(d["non_compliant"] for d in devices),
        "total_not_assessed": sum(d["not_assessed"] for d in devices),
        "total_findings": sum(d["findings_count"] for d in devices),
    }
    cumulative["avg_compliance_pct"] = (
        round(sum(d["compliance_pct"] for d in devices) / device_count, 1)
        if device_count
        else 0.0
    )
    cumulative["avg_non_compliant_pct"] = (
        round(
            sum(_pct(d["non_compliant"], d["total_controls"]) for d in devices)
            / device_count,
            1,
        )
        if device_count
        else 0.0
    )
    cumulative["avg_not_assessed_pct"] = (
        round(
            sum(_pct(d["not_assessed"], d["total_controls"]) for d in devices)
            / device_count,
            1,
        )
        if device_count
        else 0.0
    )
    cumulative["compliance_score_pct"] = _pct(
        cumulative["total_compliant"], cumulative["total_controls"]
    )

    base_meta = {}
    if devices:
        base_meta = {
            "_source": devices[0].get("source"),
            "_collected_at": devices[0].get("collected_at"),
        }

    return {
        "cumulative": cumulative,
        "devices": devices,
        "_firewall_id": "estate",
        **base_meta,
    }


def _unique_sheet_name(device, used):
    clean = "".join(
        ch for ch in device if ch.isalnum() or ch in (" ", "-", "_")
    ).strip()
    name = ("Controls - {0}".format(clean or "device"))[:31]
    base = name
    suffix = 1
    while name in used:
        name = "{0}-{1}".format(base[:28], suffix)
        suffix += 1
    used.add(name)
    return name


def _estate_excel_file(estate, output_file):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    cumulative = estate.get("cumulative") or {}
    devices = estate.get("devices") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "Estate Summary"

    title = Font(bold=True, size=15, color="1f2937")
    sub = Font(size=10, color="6b7280")
    section = Font(bold=True, size=11, color="1f2937")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="6366f1")
    alt_fill = PatternFill("solid", fgColor="f1f5f9")

    ws["A1"] = "Full Inventory Compliance Assessment"
    ws["A1"].font = title
    ws["A2"] = "All managed firewalls \u00b7 generated {0}".format(
        (estate.get("_collected_at") or "").split("T")[0]
    )
    ws["A2"].font = sub

    row = 4
    ws.cell(row=row, column=1, value="Cumulative posture").font = section
    row += 1
    summary_rows = [
        ("Average compliant", "{0}%".format(cumulative.get("avg_compliance_pct", 0.0))),
        ("Average non-compliant", "{0}%".format(cumulative.get("avg_non_compliant_pct", 0.0))),
        ("Average not assessed", "{0}%".format(cumulative.get("avg_not_assessed_pct", 0.0))),
        ("Devices assessed", cumulative.get("device_count", 0)),
        ("Total controls", cumulative.get("total_controls", 0)),
        ("Compliant controls", cumulative.get("total_compliant", 0)),
        ("Non-compliant controls", cumulative.get("total_non_compliant", 0)),
        ("Not assessed controls", cumulative.get("total_not_assessed", 0)),
        ("Open findings", cumulative.get("total_findings", 0)),
    ]
    for label, value in summary_rows:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    headers = [
        "Device",
        "Status",
        "Compliance %",
        "Compliant",
        "Non-Compliant",
        "Not Assessed",
        "Controls",
        "Findings",
        "Clone of",
        "Host IP",
    ]
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
    header_row = row
    row += 1
    for i, dev in enumerate(devices):
        values = [
            dev.get("device"),
            (dev.get("status") or "down").title(),
            "{0}%".format(dev.get("compliance_pct", 0.0)),
            dev.get("compliant", 0),
            dev.get("non_compliant", 0),
            dev.get("not_assessed", 0),
            dev.get("total_controls", 0),
            dev.get("findings_count", 0),
            dev.get("clone_of") or "\u2014",
            dev.get("host_ip") or "\u2014",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if i % 2 == 1:
                cell.fill = alt_fill
        row += 1

    for col, width in enumerate(
        [20, 12, 14, 12, 14, 14, 10, 10, 16, 18], start=1
    ):
        ws.column_dimensions[chr(64 + col)].width = width

    used = {"Estate Summary"}
    for dev in devices:
        device = dev.get("device") or "device"
        sheet = wb.create_sheet(title=_unique_sheet_name(device, used))
        r = 1
        sheet.cell(
            row=r, column=1, value="Compliance Controls \u2014 {0}".format(device)
        ).font = Font(bold=True, size=12)
        r += 2
        for col, label in enumerate(
            ["Control", "Status", "Risk", "Metric", "Observed", "Expected"],
            start=1,
        ):
            cell = sheet.cell(row=r, column=col, value=label)
            cell.font = header_font
            cell.fill = header_fill
        r += 1
        for result in dev.get("assessment") or []:
            expected = result.get("expected")
            if isinstance(expected, list):
                expected = ", ".join(str(item) for item in expected)
            values = [
                result.get("control"),
                result.get("status"),
                result.get("risk"),
                result.get("metric"),
                result.get("observed"),
                expected,
            ]
            for col, value in enumerate(values, start=1):
                sheet.cell(row=r, column=col, value=value)
            r += 1
        r += 1
        sheet.cell(
            row=r, column=1, value="Findings and remediation \u2014 {0}".format(device)
        ).font = Font(bold=True, size=11)
        r += 1
        for col, label in enumerate(
            ["Control", "Risk", "Finding", "Remediation"], start=1
        ):
            cell = sheet.cell(row=r, column=col, value=label)
            cell.font = header_font
            cell.fill = header_fill
        r += 1
        for finding in dev.get("findings") or []:
            values = [
                finding.get("control"),
                finding.get("risk"),
                finding.get("finding"),
                finding.get("remediation"),
            ]
            for col, value in enumerate(values, start=1):
                sheet.cell(row=r, column=col, value=value).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
            r += 1
        for col, width in enumerate([12, 10, 12, 60, 60, 40], start=1):
            letter = chr(64 + col)
            sheet.column_dimensions[letter].width = width
        sheet.freeze_panes = "A3"

    os.makedirs(EXCEL_DIR, exist_ok=True)
    wb.save(output_file)
    return output_file


def get_estate_excel(force=False):
    """Generate the aggregated full-inventory workbook.

    Produces an ``.xlsx`` with an estate summary tab (cumulative averages and
    a per-device table) plus a controls + findings tab for every managed
    device in the inventory.
    """
    estate = get_estate_assessment(force=force)

    month = timeutil.ist_now().strftime("%b_%Y")
    filename = "Full_Inventory_Assessment_Workbook_{0}.xlsx".format(month)
    output_file = os.path.join(EXCEL_DIR, filename)

    _estate_excel_file(estate, output_file)

    return {
        "status": "SUCCESS",
        "message": "Full inventory workbook generated successfully.",
        "summary": estate.get("cumulative") or {},
        "local_file": output_file,
        "download_url": "reports/{0}".format(filename),
        "_source": estate.get("_source", "sample"),
        "firewall_id": "estate",
    }
