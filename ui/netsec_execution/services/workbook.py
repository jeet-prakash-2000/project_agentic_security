"""Excel workbook parsing and template generation for netsec playbooks.

Workbooks are ordinary ``.xlsx`` files (built with openpyxl). Every sheet is a
table whose first row is the column header; each following row describes one
firewall change. An ``Action`` column drives the operation per row:

* ``create`` (or add/new/set)  - add the object / rule
* ``update`` (or edit)         - replace the object / rule
* ``delete`` (or remove)       - remove the object / rule

``build_template`` writes a ready-to-fill workbook containing one sheet per
playbook plus an Instructions sheet.
"""

import logging

from openpyxl import Workbook, load_workbook

logger = logging.getLogger("netsec.workbook")

INSTRUCTION_SHEET = "Instructions"


def _cell_value(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_workbook(path):
    """Parse ``path`` into ``{sheet_title: {"headers": [...], "rows": [...]}}``.

    Only data rows are returned; entirely empty rows are dropped. The
    Instructions sheet (if present) is skipped.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    try:
        for ws in wb.worksheets:
            if ws.title.strip().lower() == INSTRUCTION_SHEET.lower():
                continue
            rows = ws.iter_rows(values_only=True)
            headers = None
            parsed = []
            for values in rows:
                cells = [_cell_value(v) for v in (values or ())]
                if headers is None:
                    headers = [h for h in cells]
                    if not any(headers):
                        headers = None
                    continue
                if not any(cells):
                    continue
                record = {}
                for idx, header in enumerate(headers or []):
                    value = cells[idx] if idx < len(cells) else ""
                    if header:
                        record[header] = value
                parsed.append(record)
            if headers and any(headers):
                sheets[ws.title] = {
                    "headers": [h for h in headers if h],
                    "rows": parsed,
                    "row_count": len(parsed),
                }
    finally:
        wb.close()
    return sheets


def summarize(path):
    """Human/UI friendly summary of an uploaded workbook."""
    sheets = read_workbook(path)
    result = []
    for title, data in sheets.items():
        headers = data["headers"] or []
        result.append(
            {
                "sheet": title,
                "rows": data["row_count"],
                "columns": len(headers),
                "first_columns": headers[:8],
            }
        )
    return {"sheets": result, "total_rows": sum(r["rows"] for r in result)}


def build_template(path, playbooks):
    """Write a fill-in workbook for ``playbooks`` to ``path``."""
    wb = Workbook()
    ws_help = wb.active
    ws_help.title = INSTRUCTION_SHEET
    help_lines = [
        ["LTM NetSec Execution Agent - Playbook workbook"],
        [],
        ["Fill in one sheet per playbook. The first row is the column header;"],
        ["each row below it describes one firewall change."],
        [],
        ["Action column values:", "  create  - add the object/rule"],
        ["                        update  - replace the object/rule"],
        ["                        delete  - remove the object/rule"],
        [],
        ["Lists (tags, members, zones, rules ...) use commas, semicolons or"],
        ["new lines to separate entries, e.g. web, db"],
        [],
        ["Empty rows are ignored. Running a playbook is a DRY RUN by default,"],
        ["so nothing is sent to the firewall until the platform is configured"],
        ["with NETSEC_FW_DRY_RUN=0. In apply mode the run automatically"],
        ["commits the changes to the running configuration at the end."],
    ]
    for line in help_lines:
        ws_help.append(line)
    ws_help.column_dimensions["A"].width = 110

    for playbook in playbooks:
        sheet = wb.create_sheet(title=playbook["sheet"])
        sheet.append(playbook["columns"])
        for example in playbook.get("examples", []):
            sheet.append([example.get(col, "") for col in playbook["columns"]])
        for col in playbook.get("column_widths", {}):
            sheet.column_dimensions[col].width = playbook["column_widths"][col]
    wb.save(path)
    logger.info("Wrote playbook template to %s (%s playbooks).", path, len(playbooks))
